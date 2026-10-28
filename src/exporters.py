# -*- coding: utf-8 -*-
"""导出后端（离线可用）：Markdown → Word(.docx) / EPUB 3，以及 PDF 书签后处理。

设计原则
--------
1. **不强依赖外部工具**：本机装了 pandoc 时优先调用（转换质量最好）；
   未安装则用内置实现（python-docx / 标准库 zipfile）保证绿色便携版开箱可用。
2. **诚实的降级**：`` ```math `` 公式块与 `` ```mermaid `` 图表块无法在离线环境
   渲染成位图，导出时降级为「源码 + 标注」，而不是静默丢失。
3. 图片按 Markdown 中的相对路径相对**文档所在目录**解析；找不到则跳过该图。
"""
import base64
import html as _html
import logging
import os
import re
import shutil
import subprocess
import uuid
import zipfile
import mdutil   # v3.45（BOOKLINK-E）：锚点 slug 与编辑器同源
from i18n import T   # v3.49（I18N）：导出产物中的中文文案随界面语言

_log = logging.getLogger("mdEditor.export")

# ---------------------------------------------------------------- Markdown 解析

_HEAD = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_HR = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")
_FENCE_OPEN = re.compile(r"^\s{0,3}(`{3,}|~{3,})\s*([^\s`]*)\s*$")
_QUOTE = re.compile(r"^\s{0,3}>\s?(.*)$")
_ULI = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_OLI = re.compile(r"^(\s*)\d{1,9}[.)]\s+(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")

# 行内元素：图片 / 行内代码 / 粗体 / 斜体 / 删除线 / 链接 / 行内公式
# （v3.30.11 MATH-2：行内 $...$ 识别后降级为去 $ 纯文本，DOCX/EPUB 不再出现
#   裸 $ 标记；护栏同 editor —— 内容首尾非空白、闭合 $ 后不跟字母/数字/$，
#   避免「价格 $5，再 $6」误判）
_INLINE = re.compile(
    r"(!\[[^\]]*\]\([^)]*\))"
    r"|(\[\^[^\]\s]+\])"
    r"|(`[^`]+`)"
    r"|(\*\*[^*]+\*\*|__[^_]+__)"
    r"|(\*[^*]+\*|_[^_]+_)"
    r"|(~~[^~]+~~)"
    r"|(\[[^\]]+\]\([^)]*\))"
    r"|(\[\[[^\]]+\]\])"          # v3.45（BOOKLINK-E）：双链 / 本书链接 [[Target#anchor|alias]]
    r"|(\$(?!\s)[^$\n]+?(?<!\s)\$(?![\w$]))"
)
_IMG_ONLY = re.compile(r"^!\[([^\]]*)\]\(\s*([^)\s]+)(?:\s+\"[^\"]*\")?\s*\)$")
# v3.30.18（FN-2 / DL-5）：脚注引用、脚注定义行、定义列表定义行
_FNREF = re.compile(r"\[\^([^\]\s]+)\]")
_FNDEF = re.compile(r"^\s{0,3}\[\^([^\]\s]+)\]:\s*(.*)$")
_DL_DEF = re.compile(r"^\s{0,3}:\s+(.*)$")


def _split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse_blocks(md):
    """把 Markdown 切成结构化块列表（导出用，够用即可）。"""
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        m = _FENCE_OPEN.match(line)
        if m:
            mark, lang = m.group(1)[0], (m.group(2) or "").strip().lower()
            i += 1
            buf = []
            while i < n and not re.match(r"^\s{0,3}" + re.escape(mark) + r"{3,}\s*$", lines[i]):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过结束围栏
            blocks.append({"kind": "fence", "lang": lang, "text": "\n".join(buf)})
            continue

        m = _FNDEF.match(line)
        if m:
            # v3.30.18（FN-2）：脚注定义独立成块
            blocks.append({"kind": "fndef", "id": m.group(1), "text": m.group(2)})
            i += 1
            continue

        m = _HEAD.match(line)
        if m:
            blocks.append({"kind": "heading", "level": len(m.group(1)), "text": m.group(2)})
            i += 1
            continue

        if _HR.match(line):
            blocks.append({"kind": "hr"})
            i += 1
            continue

        # v3.30.19（MATH-6）：$$…$$ 块级公式 → 按公式源码块降级导出（渲染由
        # 编辑器 / pandoc 通道负责；内置导出与 ```math 同路径）
        m6 = re.match(r"^\s{0,3}\$\$(.*)$", line)
        if m6:
            first6 = m6.group(1)
            i += 1
            if first6.strip().endswith("$$"):
                blocks.append({"kind": "fence", "lang": "math",
                               "text": first6.strip()[:-2].strip()})
                continue
            buf6 = []
            if first6.strip():
                buf6.append(first6)
            while i < n and "$$" not in lines[i]:
                buf6.append(lines[i])
                i += 1
            if i < n:
                buf6.append(re.sub(r"\$\$.*$", "", lines[i]))
                i += 1
            blocks.append({"kind": "fence", "lang": "math",
                           "text": "\n".join(buf6).strip()})
            continue

        if "|" in line and i + 1 < n and _TABLE_SEP.match(lines[i + 1]):
            header = _split_row(line)
            aligns = _split_row(lines[i + 1])
            i += 2
            rows = []
            while i < n and lines[i].strip() and "|" in lines[i]:
                rows.append(_split_row(lines[i]))
                i += 1
            blocks.append({"kind": "table", "header": header, "aligns": aligns, "rows": rows})
            continue

        if _QUOTE.match(line):
            buf = []
            while i < n and _QUOTE.match(lines[i]):
                buf.append(_QUOTE.match(lines[i]).group(1))
                i += 1
            blocks.append({"kind": "quote", "text": "\n".join(buf).strip()})
            continue

        mu, mo = _ULI.match(line), _OLI.match(line)
        if mu or mo:
            ordered = bool(mo)
            items = []
            while i < n:
                a, b = _ULI.match(lines[i]), _OLI.match(lines[i])
                if not (a or b) or bool(b) != ordered:
                    break
                mm = b or a
                items.append({"indent": len(mm.group(1)) // 2, "text": mm.group(2)})
                i += 1
            blocks.append({"kind": "list", "ordered": ordered, "items": items})
            continue

        # 普通段落：合并连续非空行，遇到其它块起始则停止
        # （v3.30.18（DL-5）：": 定义" 行视为定义列表组成部分，段落收集在此停住）
        buf = [line]
        i += 1
        while i < n and lines[i].strip() and not any(r.match(lines[i]) for r in (
                _HEAD, _FENCE_OPEN, _QUOTE, _ULI, _OLI, _HR, _DL_DEF)):
            buf.append(lines[i])
            i += 1
        if i < n and _DL_DEF.match(lines[i]) and buf:
            # v3.30.18（DL-5）：术语 + ": 定义"... → 定义列表块
            defs = []
            while i < n and _DL_DEF.match(lines[i]):
                defs.append(_DL_DEF.match(lines[i]).group(1).strip())
                i += 1
            blocks.append({"kind": "deflist",
                           "term": "\n".join(buf).strip(), "defs": defs})
            continue
        blocks.append({"kind": "para", "text": "\n".join(buf)})
    return blocks


def parse_inline(text):
    """行内元素解析 → [{text, bold, italic, code, strike, link, image}]。"""
    out, pos = [], 0
    for m in _INLINE.finditer(text or ""):
        if m.start() > pos:
            out.append({"text": text[pos:m.start()]})
        tok = m.group(0)
        if m.group(1):                      # 图片
            im = _IMG_ONLY.match(tok)
            out.append({"image": im.group(2) if im else "", "alt": im.group(1) if im else ""})
        elif m.group(2):                    # 脚注引用 [^id]（v3.30.18 FN-2）
            out.append({"fnref": tok[2:-1]})
        elif m.group(3):                    # 行内代码
            out.append({"text": tok[1:-1], "code": True})
        elif m.group(4):                    # 粗体
            out.append({"text": tok[2:-2], "bold": True})
        elif m.group(5):                    # 斜体
            out.append({"text": tok[1:-1], "italic": True})
        elif m.group(6):                    # 删除线
            out.append({"text": tok[2:-2], "strike": True})
        elif m.group(7):                    # 链接
            lb = re.match(r"\[([^\]]+)\]\(([^)]*)\)", tok)
            out.append({"text": lb.group(1) if lb else tok,
                        "link": (lb.group(2) if lb else "")})
        elif m.group(8):                    # v3.45（BOOKLINK-E）：双链 / 本书链接 [[Target#anchor|alias]]
            inner = tok[2:-2].strip()
            target, _, alias = inner.partition("|")
            target = target.strip()
            anchor = ""
            if "#" in target:
                target, _, anchor = target.partition("#")
                target = target.strip()
                anchor = anchor.strip()
            out.append({"wikilink": True, "target": target,
                        "anchor": anchor, "alias": alias.strip()})
        elif m.group(9):                    # 行内公式 $...$ → 标记 math（导出优先位图化，回落纯文本）
            out.append({"math": tok[1:-1]})
        pos = m.end()
    if pos < len(text or ""):
        out.append({"text": text[pos:]})
    return out or [{"text": text or ""}]


def plain_inline(text):
    """行内元素 → 纯文本（降级、图表源码等场景用）。"""
    parts = []
    for t in parse_inline(text):
        if t.get("image"):
            parts.append(t.get("alt") or T("图片", "image"))
        elif t.get("fnref"):
            parts.append("[" + t["fnref"] + "]")
        elif t.get("math") is not None:
            parts.append(t["math"])
        else:
            parts.append(t.get("text") or "")
    return "".join(parts)


# ---------------------------------------------------------------- Word (.docx)

_CJK_FONT = "微软雅黑"
_MONO_FONT = "Consolas"


def _set_style_font(style, name, size=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)


def _run_font(run, name):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), name)


def _shade(paragraph, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


# v3.49（EXPORT-DOCX）：表格美观化辅助 —— 单元格底色 / 边距 / 垂直居中 / 列宽
def _set_cell_bg(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """单元格内边距（dxa，1/20 磅）；start/end + left/right 同时给以兼容各阅读器。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom),
                     ("start", left), ("end", right),
                     ("left", left), ("right", right)):
        e = OxmlElement("w:" + tag)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def _set_cell_valign(cell, valign="center"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), valign)
    tcPr.append(v)


# 表格分隔单元格（:---: / :---- / ----: / -----）→ 对齐方式枚举值
_TABLE_ALIGN = {0: None, 1: "center", 2: "right"}
_HEAD_SIZES = {1: 18, 2: 15, 3: 13, 4: 12}


def _add_runs(paragraph, text, base_dir, doc, allow_image=True, math_png=None):
    """把一行 Markdown 行内内容写进段落（图片单独成段）。

    v3.44.31（MATH-EXP）：math_png 为 {公式源码: PNG 路径} 映射（由 GUI 侧把已渲染
    公式位图化后传入）；行内公式命中时直接嵌 PNG（与编辑器显示一致），未命中回落纯文本。
    """
    from docx.shared import Inches, Pt
    for t in parse_inline(text):
        if t.get("math") is not None:
            # 行内公式：优先位图嵌入（I: 前缀区分行内），失败回落斜体纯文本
            tex = t["math"]
            png = (math_png or {}).get("I:" + tex.strip())
            if png and os.path.isfile(png):
                try:
                    run = paragraph.add_run()
                    run.add_picture(png, height=Pt(13))
                    continue
                except Exception as e2:
                    _log.debug("docx 行内公式嵌入失败，回落纯文本: %s", e2)
            r = paragraph.add_run(tex)
            r.italic = True
            continue
        if t.get("image"):
            if not allow_image:
                continue
            p = _add_picture(doc, t["image"], base_dir, Inches(5.6))
            if p is None:
                r = paragraph.add_run(f"[图片：{t.get('alt') or os.path.basename(t['image'])}]")
                r.italic = True
            continue
        if t.get("fnref"):
            # v3.30.18（FN-2）：脚注引用 → 上标 [id]（Word 原生脚注对象成本高，
            # 按「正文上标 [n] + 文末定义段」降级，语义不再按字面丢失）
            r2 = paragraph.add_run("[" + t["fnref"] + "]")
            r2.font.superscript = True
            continue
        r = paragraph.add_run(t.get("text") or "")
        r.bold = bool(t.get("bold"))
        r.italic = bool(t.get("italic")) or bool(t.get("link"))
        r.font.strike = bool(t.get("strike"))
        if t.get("code"):
            _run_font(r, _MONO_FONT)
        if t.get("link") and t.get("text"):
            r.font.underline = True
    return paragraph


def _add_picture(doc, rel, base_dir, width):
    """按相对路径插入图片；文件不存在返回 None。"""
    if not rel or rel.startswith(("http://", "https://", "data:")):
        return None
    path = rel if os.path.isabs(rel) else os.path.join(base_dir or "", rel.replace("/", os.sep))
    if not os.path.isfile(path):
        return None
    try:
        p = doc.add_paragraph()
        p.alignment = 1  # 居中
        p.add_run().add_picture(path, width=width)
        return p
    except Exception:
        return None


def write_docx(md, out_path, base_dir="", title=None, mmd_png=None, math_png=None):
    """Markdown → .docx。返回 (是否成功, 提示信息)。

    v3.30.19（MMD-2）：mmd_png 为 {图表源码: PNG 路径} 映射（由 GUI 侧把已渲染
    SVG 位图化后传入）；命中时图表直接嵌 PNG，未命中回落「源码 + 标注」降级。
    """
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor, Mm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
    except ImportError:
        return False, "缺少 python-docx，无法导出 Word（可安装 pandoc 或 pip install python-docx）"

    doc = docx.Document()
    # v3.49（EXPORT-DOCX）：A4 标准纸张 + 2.54cm 页边距（不再跟随系统默认 Letter）
    _sec = doc.sections[0]
    _sec.page_width = Mm(210)
    _sec.page_height = Mm(297)
    _sec.left_margin = Mm(25.4)
    _sec.right_margin = Mm(25.4)
    _sec.top_margin = Mm(25.4)
    _sec.bottom_margin = Mm(25.4)
    # 正文：微软雅黑 10.5pt、1.5 倍行距、段后 6pt，阅读更舒展
    _n = doc.styles["Normal"]
    _set_style_font(_n, _CJK_FONT, 10.5)
    _n.paragraph_format.line_spacing = 1.5
    _n.paragraph_format.space_after = Pt(6)
    for sname in ("List Bullet", "List Number", "Quote"):
        if sname in [s.name for s in doc.styles]:
            try:
                _set_style_font(doc.styles[sname], _CJK_FONT, 10.5)
            except Exception as e:   # v3.29d（B9）：记录而非静默吞掉
                _log.debug("docx 样式 %s 字体设置失败: %s", sname, e)

    for blk in parse_blocks(md):
        k = blk["kind"]
        if k == "heading":
            lvl = max(1, min(4, blk["level"]))
            h = doc.add_heading("", level=lvl)
            try:
                _set_style_font(doc.styles[f"Heading {lvl}"], _CJK_FONT)
            except Exception as e:   # v3.29d（B9）：记录而非静默吞掉
                _log.debug("docx 标题样式 Heading %s 字体设置失败: %s", lvl, e)
            _add_runs(h, blk["text"], base_dir, doc, allow_image=False, math_png=math_png)
            for r in h.runs:
                r.font.size = Pt(_HEAD_SIZES[lvl])
                r.bold = True
                r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(6)
        elif k == "para":
            p = doc.add_paragraph()
            _add_runs(p, blk["text"], base_dir, doc, math_png=math_png)
        elif k == "list":
            style = "List Number" if blk["ordered"] else "List Bullet"
            for it in blk["items"]:
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph()
                if it["indent"]:
                    p.paragraph_format.left_indent = Pt(18 * (it["indent"] + 1))
                _add_runs(p, it["text"], base_dir, doc, math_png=math_png)
        elif k == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _add_runs(p, blk["text"], base_dir, doc, math_png=math_png)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        elif k == "fence":
            lang, code = blk["lang"], blk["text"]
            if lang == "mermaid" and mmd_png and mmd_png.get(code) \
                    and os.path.isfile(mmd_png[code]):
                # v3.30.19（MMD-2）：已位图化的图表直接嵌 PNG（2x 渲染，白底）
                try:
                    p = doc.add_paragraph()
                    p.alignment = 1  # 居中
                    p.add_run().add_picture(mmd_png[code], width=Inches(5.6))
                    continue
                except Exception as e2:
                    _log.debug("docx 图表 PNG 嵌入失败，回落源码降级: %s", e2)
            if lang == "math":
                # v3.44.31（MATH-EXP）：已位图化的公式（B: 前缀区分块级）直接嵌 PNG
                png = (math_png or {}).get("B:" + (code or "").strip())
                if png and os.path.isfile(png):
                    try:
                        p = doc.add_paragraph()
                        p.alignment = 1  # 居中
                        p.add_run().add_picture(png, width=Inches(5.6))
                        continue
                    except Exception as e2:
                        _log.debug("docx 公式 PNG 嵌入失败，回落源码降级: %s", e2)
                tip = doc.add_paragraph()
                _add_runs(tip, "［公式 · LaTeX 源码，渲染失败回落］", base_dir, doc, math_png=math_png)
                for r in tip.runs:
                    r.italic = True
                    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            elif lang == "mermaid":
                tip = doc.add_paragraph()
                _add_runs(tip, "［图表 · Mermaid 源码，请在 mdEditor 中查看渲染效果］", base_dir, doc, math_png=math_png)
                for r in tip.runs:
                    r.italic = True
                    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            for ln in (code or "").split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                _shade(p, "F4F5F7")
                r = p.add_run(ln if ln else " ")
                _run_font(r, _MONO_FONT)
                r.font.size = Pt(9.5)
        elif k == "table":
            header = blk["header"] or [""]
            aligns = blk.get("aligns") or []
            cols = len(header)
            table = doc.add_table(rows=1, cols=cols)
            try:
                table.style = "Table Grid"
            except KeyError:
                pass
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            # v3.49（EXPORT-DOCX）：A4 内容区约 6.2in，列宽均匀分配
            from docx.shared import Inches as _Inches
            col_w = max(0.6, 6.2 / cols)

            def _align_for(idx):
                s = (aligns[idx] if idx < len(aligns) else "").strip()
                if s.startswith(":") and s.endswith(":"):
                    return WD_ALIGN_PARAGRAPH.CENTER
                if s.endswith(":"):
                    return WD_ALIGN_PARAGRAPH.RIGHT
                return WD_ALIGN_PARAGRAPH.LEFT

            for c, txt in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                _set_cell_bg(cell, "EAF1FB")
                _set_cell_valign(cell, "center")
                _set_cell_margins(cell)
                p = cell.paragraphs[0]
                p.alignment = _align_for(c)
                _add_runs(p, txt, base_dir, doc, allow_image=False, math_png=math_png)
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
                cell.width = _Inches(col_w)
            for ri, row in enumerate(blk["rows"]):
                cells = table.add_row().cells
                for c in range(cols):
                    cell = cells[c]
                    cell.text = ""
                    _set_cell_valign(cell, "center")
                    _set_cell_margins(cell)
                    if ri % 2 == 1:
                        _set_cell_bg(cell, "F5F8FC")   # 斑马纹，提升可读性
                    p = cell.paragraphs[0]
                    p.alignment = _align_for(c)
                    _add_runs(p, row[c] if c < len(row) else "",
                              base_dir, doc, allow_image=False, math_png=math_png)
                    cell.width = _Inches(col_w)
        elif k == "fndef":
            # v3.30.18（FN-2）：脚注定义 → 文末标签段
            p = doc.add_paragraph()
            rlab = p.add_run("[" + blk["id"] + "]")
            rlab.bold = True
            p.add_run(" " + (blk["text"] or ""))
            for r in p.runs:
                r.font.size = Pt(9.5)
        elif k == "deflist":
            # v3.30.18（DL-5）：定义列表 → 术语加粗段 + 定义缩进段
            pt = doc.add_paragraph()
            rt = pt.add_run(plain_inline(blk.get("term") or ""))
            rt.bold = True
            for d in blk.get("defs") or []:
                pd = doc.add_paragraph()
                pd.paragraph_format.left_indent = Pt(18)
                _add_runs(pd, d, base_dir, doc, math_png=math_png)
        elif k == "hr":
            p = doc.add_paragraph()
            _shade(p, "DDDDDD")

    try:
        doc.save(out_path)
    except Exception as e:
        return False, f"写入 docx 失败：{e}"
    return True, f"已导出 Word · {out_path}"


# ---------------------------------------------------------------- EPUB 3

_EPUB_CSS = """body{font-family:"Microsoft YaHei","PingFang SC",serif;line-height:1.75;
margin:0 5%;word-wrap:break-word}
h1,h2,h3,h4{line-height:1.35;margin:1.4em 0 .6em}
h1{font-size:1.6em;border-bottom:1px solid #ddd;padding-bottom:.3em}
pre{background:#f5f6f8;padding:.8em;border-radius:6px;overflow-x:auto;font-size:.86em;
white-space:pre-wrap;word-break:break-all}
code{font-family:Consolas,Menlo,monospace;background:#f2f3f5;border-radius:3px;padding:0 .25em}
pre code{background:none;padding:0}
blockquote{border-left:4px solid #d0d7de;margin:1em 0;padding:.2em 1em;color:#57606a}
table{border-collapse:collapse;width:100%;margin:1em 0}
th,td{border:1px solid #d0d7de;padding:.4em .6em;text-align:left}
th{background:#f5f6f8}
img{max-width:100%;height:auto}
hr{border:none;border-top:1px solid __HR_BORDER__;margin:1.6em 0}
.tip{color:#8a8a8a;font-style:italic}
.imath{vertical-align:middle}
.mblock{display:block;margin:0 auto}
sup.fnref a{text-decoration:none;font-size:.75em}
p.fndef{font-size:.88em;color:#444;margin:.3em 0}
dl dt{font-weight:700;margin-top:.6em}
dl dd{margin:0 0 .4em 1.6em}
"""

_ESC = {"&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;"}


def _xesc(s):
    return "".join(_ESC.get(c, c) for c in str(s or ""))


def _xinline(text, base_dir="", math_png=None, images=None, link_map=None,
             file_map=None):
    out = []
    for t in parse_inline(text):
        if t.get("image"):
            out.append(f'<img src="images/{_xesc(_img_name(t["image"]))}" '
                       f'alt="{_xesc(t.get("alt") or "")}"/>')
            continue
        if t.get("wikilink"):
            # v3.45（BOOKLINK-E）：[[Target#anchor|alias]] → <a href="page.html#slug">
            target, anchor, alias = t["target"], t["anchor"], t.get("alias") or ""
            href = None
            if link_map and target:
                href = link_map.get(target) or link_map.get(target.lower())
                if href and anchor:
                    href = href + "#" + anchor
            if href:
                out.append('<a href="%s">%s</a>' % (_xesc(href), _xesc(alias or target)))
            else:
                # 无法解析（非书籍 / 目标缺失）：回落纯文本，避免导出死链
                out.append(_xesc(alias or target))
            continue
        if t.get("fnref"):
            # v3.30.18（FN-2）：脚注引用 → 上标锚链接（跳到文末定义）
            fid = _xesc(t["fnref"])
            out.append(f'<sup class="fnref"><a href="#fn-{fid}">[{fid}]</a></sup>')
            continue
        if t.get("math") is not None:
            # v3.44.31（MATH-EXP）：行内公式优先位图嵌入（I: 前缀区分行内），
            # 未命中回落纯文本
            tex = t["math"]
            png = (math_png or {}).get("I:" + tex.strip())
            if png:
                img_name = f"math_{uuid.uuid4().hex[:8]}.png"
                if images is not None:
                    images[img_name] = png
                out.append(f'<img class="imath" src="images/{img_name}" '
                           f'alt="{_xesc(tex)}"/>')
            else:
                out.append(_xesc(tex))
            continue
        s = _xesc(t.get("text") or "")
        if not s:
            continue
        if t.get("code"):
            s = f"<code>{s}</code>"
        if t.get("bold"):
            s = f"<strong>{s}</strong>"
        if t.get("italic"):
            s = f"<em>{s}</em>"
        if t.get("strike"):
            s = f"<del>{s}</del>"
        if t.get("link"):
            # v3.45.3（EXPORT-LINK）：本地 .md 链接改写为 .html 目标
            # （命中章节映射走页面文件，否则换后缀）
            href = _rewrite_export_link(t["link"], base_dir, file_map) \
                if base_dir else t["link"]
            s = f'<a href="{_xesc(href)}">{s}</a>'
        out.append(s)
    return "".join(out)

def _img_name(rel):
    return os.path.basename(str(rel).replace("\\", "/"))


# v3.45.3（EXPORT-LINK）：站点 / 导出物内的本地 .md 链接改写
_LINK_ABS_PREFIX = ("http://", "https://", "mailto:", "data:", "blob:",
                    "file:", "#", "/", "//")
_MD_EXT_RE = re.compile(r"\.(md|markdown|mdown)$", re.I)


def _rewrite_export_link(url, base_dir, file_map=None):
    """把导出物里的本地 .md 链接改写为 .html 目标。

    - 外部 / 锚点 / data / 绝对路径等原样返回；
    - 命中本书章节（file_map: {abs_path_lower: 页面文件}）→ 对应页面
      （NN-xxx.html），保留 #fragment；
    - 其它本地 .md → 后缀换为 .html（页面未生成时也指向正确类型）。
    返回改写后的 href（无需改写时原样返回）。"""
    u = str(url or "").strip()
    if not u:
        return url
    low = u.lower()
    for p in _LINK_ABS_PREFIX:
        if low.startswith(p):
            return url
    path, sep, frag = u.partition("#")
    if not _MD_EXT_RE.search(path):
        return url
    tail = (sep + frag) if sep else ""
    ap = os.path.normpath(os.path.join(base_dir or ".",
                                       path.replace("\\", "/")))
    key = os.path.normcase(ap).lower()
    if file_map and key in file_map:
        return file_map[key] + tail
    root, _ext = os.path.splitext(path)
    return root + ".html" + tail


def _blocks_to_xhtml(blocks, base_dir, images, mmd_png=None, math_png=None,
                     link_map=None, file_map=None):
    """块列表 → XHTML 片段；顺带收集需要打包进 EPUB 的图片。

    v3.30.19（MMD-2）：mmd_png 命中的图表直接以 PNG 内嵌（随 EPUB 打包）。
    v3.44.31（MATH-EXP）：math_png 命中的公式（块级 B: 前缀）同样以 PNG 内嵌；
    行内公式经 _xinline 内嵌。未命中回落源码降级。
    v3.45（BOOKLINK-E）：link_map 为 {标题: 页面文件} 时，[[双链]] 转为跨页
    <a href="page.html#slug">；标题（h1-h4）按章内去重写入 id（与编辑器同源）。
    v3.45.3（EXPORT-LINK）：file_map 为 {章节abs路径: 页面文件} 时，正文本地
    .md 链接改写为对应 .html 页面（保留 #fragment）；非章节 .md 统一换后缀。
    """
    out = []
    _hseen = {}   # v3.45：章内标题锚点去重（仅 h1-h4，与选择器/编辑器一致）
    for blk in blocks:
        k = blk["kind"]
        if k == "heading":
            lvl = max(1, min(6, blk["level"]))
            _id = ""
            if lvl <= 4:
                _base = mdutil.slugify(blk["text"]) or "section"
                if _base in _hseen:
                    _hseen[_base] += 1
                    _slug = "%s-%d" % (_base, _hseen[_base])
                else:
                    _hseen[_base] = 1
                    _slug = _base
                _id = ' id="%s"' % _xesc(_slug)
            out.append(f"<h{lvl}{_id}>{_xinline(blk['text'], math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</h{lvl}>")
        elif k == "para":
            out.append(f"<p>{_xinline(blk['text'], math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</p>")
        elif k == "list":
            tag = "ol" if blk["ordered"] else "ul"
            out.append(f"<{tag}>")
            for it in blk["items"]:
                out.append(f"<li>{_xinline(it['text'], math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</li>")
            out.append(f"</{tag}>")
        elif k == "quote":
            out.append(f"<blockquote>{_xinline(blk['text'], math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</blockquote>")
        elif k == "hr":
            out.append("<hr/>")
        elif k == "table":
            out.append("<table><thead><tr>")
            for c in blk["header"]:
                out.append(f"<th>{_xinline(c, math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</th>")
            out.append("</tr></thead><tbody>")
            for row in blk["rows"]:
                out.append("<tr>")
                for c in range(len(blk["header"])):
                    out.append(f"<td>{_xinline(row[c] if c < len(row) else '', math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</td>")
                out.append("</tr>")
            out.append("</tbody></table>")
        elif k == "fndef":
            # v3.30.18（FN-2）：脚注定义 → 带锚点的定义段
            fid = _xesc(blk["id"])
            out.append(f'<p class="fndef" id="fn-{fid}">'
                       f'<strong>[{fid}]</strong> {_xinline(blk.get("text") or "", math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</p>')
        elif k == "deflist":
            # v3.30.18（DL-5）：定义列表 → 原生 <dl>（阅读器保留语义）
            out.append("<dl>")
            out.append(f"<dt>{_xinline(blk.get('term') or '', math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</dt>")
            for d in blk.get("defs") or []:
                out.append(f"<dd>{_xinline(d, math_png=math_png, images=images, link_map=link_map, base_dir=base_dir, file_map=file_map)}</dd>")
            out.append("</dl>")
        elif k == "fence":
            lang, code = blk["lang"], blk["text"] or ""
            if lang == "mermaid" and mmd_png and mmd_png.get(code) \
                    and os.path.isfile(mmd_png[code]):
                # v3.30.19（MMD-2）：图表 PNG 随 EPUB 打包内嵌
                img_name = f"mmd-{len(images) + 1}.png"
                images[img_name] = mmd_png[code]
                out.append(f'<p><img src="images/{img_name}" alt="{T("图表", "chart")}"/></p>')
                continue
            if lang == "math":
                # v3.44.31（MATH-EXP）：已位图化的公式（B: 前缀区分块级）直接内嵌
                png = (math_png or {}).get("B:" + code.strip())
                if png and os.path.isfile(png):
                    img_name = f"math_{uuid.uuid4().hex[:8]}.png"
                    images[img_name] = png
                    out.append(f'<p><img class="mblock" src="images/{img_name}" alt="{T("公式", "formula")}"/></p>')
                    continue
            if lang in ("mermaid", "math"):
                label = "图表（Mermaid 源码）" if lang == "mermaid" else "公式（LaTeX 源码）"
                out.append(f'<p class="tip">［{label}］</p>')
            out.append(f"<pre><code>{_xesc(code)}</code></pre>")
        # 收集图片
        if k in ("para", "quote", "heading"):
            for t in parse_inline(blk.get("text") or ""):
                if t.get("image"):
                    rel = str(t["image"])
                    if rel.startswith(("http://", "https://", "data:")):
                        continue
                    path = rel if os.path.isabs(rel) else os.path.join(
                        base_dir or "", rel.replace("/", os.sep))
                    if os.path.isfile(path):
                        images.setdefault(_img_name(rel), path)
    return "\n".join(out)


def write_epub(md, out_path, base_dir="", title=None, author="mdEditor", mmd_png=None,
               border_color="#ddd", math_png=None):
    """Markdown → EPUB 3。返回 (是否成功, 提示信息)。（MMD-2：mmd_png 见 write_docx）"""
    blocks = parse_blocks(md)
    # 按一级标题分章（无 h1 则整篇一章）
    chapters, cur = [], []
    for b in blocks:
        if b["kind"] == "heading" and b["level"] == 1 and cur:
            chapters.append(cur)
            cur = []
        cur.append(b)
    if cur:
        chapters.append(cur)
    if not chapters:
        chapters = [[]]

    title = title or "未命名文档"
    images = {}
    bodies, toc = [], []
    for i, ch in enumerate(chapters, 1):
        head = next((b["text"] for b in ch
                     if b["kind"] == "heading" and b["level"] == 1), None)
        if i == 1 and len(chapters) == 1:
            head = head or title
        cname = f"chapter{i}.xhtml"
        ctitle = plain_inline(head) if head else f"第 {i} 章"
        toc.append((ctitle, cname))
        body = _blocks_to_xhtml(ch, base_dir, images, mmd_png=mmd_png, math_png=math_png)
        bodies.append((cname, ctitle, body))

    book_id = "urn:uuid:" + str(uuid.uuid4())
    manifest, spine, nav_items = [], [], []
    manifest.append('<item id="css" href="style.css" media-type="text/css"/>')
    manifest.append('<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" '
                    'properties="nav"/>')
    for i, (cname, ctitle, _) in enumerate(bodies, 1):
        manifest.append(f'<item id="c{i}" href="{cname}" media-type="application/xhtml+xml"/>')
        spine.append(f'<itemref idref="c{i}"/>')
        nav_items.append(f'<li><a href="{cname}">{_xesc(ctitle)}</a></li>')
    for n, rel in enumerate(sorted(images), 1):
        ext = os.path.splitext(rel)[1].lower().lstrip(".") or "png"
        mt = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
              "gif": "image/gif", "svg": "image/svg+xml", "webp": "image/webp"}.get(ext, "image/png")
        manifest.append(f'<item id="img{n}" href="images/{_xesc(rel)}" media-type="{mt}"/>')

    opf = f"""<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="bookid">{book_id}</dc:identifier>
    <dc:title>{_xesc(title)}</dc:title>
    <dc:language>zh-CN</dc:language>
    <dc:creator>{_xesc(author)}</dc:creator>
    <meta property="dcterms:modified">2026-01-01T00:00:00Z</meta>
  </metadata>
  <manifest>
    {chr(10).join('  ' + m for m in manifest)}
  </manifest>
  <spine>
    {chr(10).join('  ' + s for s in spine)}
  </spine>
</package>
"""
    nav = f"""<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">
<head><title>目录</title><link rel="stylesheet" href="style.css"/></head>
<body><nav epub:type="toc" id="toc"><h1>目录</h1><ol>
{chr(10).join(nav_items)}
</ol></nav></body></html>
"""
    container = """<?xml version="1.0" encoding="utf-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf"
    media-type="application/oebps-package+xml"/></rootfiles>
</container>
"""
    try:
        with zipfile.ZipFile(out_path, "w") as z:
            # mimetype 必须第一个写入且不压缩
            z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip",
                       compress_type=zipfile.ZIP_STORED)
            z.writestr("META-INF/container.xml", container, zipfile.ZIP_DEFLATED)
            z.writestr("OEBPS/content.opf", opf, zipfile.ZIP_DEFLATED)
            z.writestr("OEBPS/nav.xhtml", nav, zipfile.ZIP_DEFLATED)
            # v3.38（HR-8②）：hr 配色按当前主题 --border 注入（默认回落 #ddd）
            z.writestr("OEBPS/style.css",
                       _EPUB_CSS.replace("__HR_BORDER__", border_color or "#ddd"),
                       zipfile.ZIP_DEFLATED)
            for cname, ctitle, body in bodies:
                page = (f'<?xml version="1.0" encoding="utf-8"?>\n'
                        f'<html xmlns="http://www.w3.org/1999/xhtml" lang="zh-CN">\n'
                        f'<head><title>{_xesc(ctitle)}</title>'
                        f'<link rel="stylesheet" href="style.css"/></head>\n'
                        f'<body>\n{body}\n</body></html>\n')
                z.writestr(f"OEBPS/{cname}", page, zipfile.ZIP_DEFLATED)
            for rel, path in images.items():
                z.write(path, f"OEBPS/images/{rel}")
    except Exception as e:
        return False, f"写入 EPUB 失败：{e}"
    return True, f"已导出 EPUB（{len(bodies)} 章节 / {len(images)} 图）· {out_path}"


# ---------------------------------------------------------------- PDF 书签

def add_pdf_bookmarks(pdf_path, headings):
    """按标题层级写入 PDF 书签。

    printToPdf 无法直接产出大纲，这里用「标题文本检索定位页码」的启发式方案：
    依次在各页文本中查找标题文字，命中即把书签指向该页并保持层级嵌套。
    返回 (写入条数, 说明)。
    """
    heads = [h for h in (headings or []) if str(h.get("text") or "").strip()]
    if not heads:
        return 0, "无标题，未写入书签"
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError:
        return 0, "缺少 pypdf，未写入 PDF 书签"

    try:
        reader = PdfReader(pdf_path)
        pages = []
        for p in reader.pages:
            try:
                pages.append(re.sub(r"\s+", "", p.extract_text() or ""))
            except Exception:
                pages.append("")
        writer = PdfWriter(clone_from=pdf_path)
        parents, added, cursor = {}, 0, 0
        for h in heads:
            key = re.sub(r"\s+", "", str(h.get("text") or ""))
            page_no = None
            for idx in list(range(cursor, len(pages))) + list(range(0, cursor)):
                if key and key in pages[idx]:
                    page_no = idx
                    break
            if page_no is None:
                continue
            cursor = page_no
            lvl = max(1, int(h.get("level") or 1))
            parent = None
            for pl in range(lvl - 1, 0, -1):
                if pl in parents:
                    parent = parents[pl]
                    break
            try:
                item = writer.add_outline_item(str(h.get("text")), page_no, parent=parent)
            except Exception as e:   # v3.29d（B9）：记录而非静默吞掉
                _log.debug("PDF 书签添加失败（%s）: %s", h.get("text"), e)
                continue
            parents[lvl] = item
            for k in [k for k in parents if k > lvl]:
                del parents[k]
            added += 1
        if added == 0:
            return 0, "未能在 PDF 中定位标题页码，未写入书签"
        with open(pdf_path, "wb") as f:
            writer.write(f)
        return added, f"已写入 {added} 条书签"
    except Exception as e:
        return 0, f"PDF 书签写入失败：{e}"


def merge_book_pdfs(chapter_pdfs, out_path, bookmarks=None):
    """合并若干章节 PDF（有序），并在指定页偏移写入层级书签。

    chapter_pdfs: [(pdf_path, page_count), ...]（文档顺序）。
    bookmarks: [(title, level, page_offset), ...]，page_offset 为 0 基页号。
    返回 (写入书签数, 说明)。
    """
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    offsets = []
    cur = 0
    total_pages = 0
    for (p, _c) in chapter_pdfs:
        try:
            reader = PdfReader(p)
        except Exception as e:
            _log.warning("章节 PDF 读取失败，跳过 %s: %s", p, e)
            continue
        offsets.append(cur)
        for pg in reader.pages:
            writer.add_page(pg)
        cur += len(reader.pages)
        total_pages += len(reader.pages)
    added = 0
    if bookmarks:
        parents = {}
        for (title, level, off) in bookmarks:
            if off < 0 or off >= total_pages:
                continue
            parent = None
            for pl in range(level - 1, 0, -1):
                if pl in parents:
                    parent = parents[pl]
                    break
            try:
                item = writer.add_outline_item(title, off, parent=parent)
            except Exception as e:
                _log.debug("书签添加失败（%s）: %s", title, e)
                continue
            parents[level] = item
            for k in [k for k in parents if k > level]:
                del parents[k]
            added += 1
    with open(out_path, "wb") as f:
        writer.write(f)
    return added, "已合并 %d 章 / %d 页，写入 %d 条书签" % (
        len(chapter_pdfs), total_pages, added)


def stamp_page_numbers(pdf_path, start=1, size=9, color=(120, 120, 120)):
    """给每页底部居中盖静音灰页码（依赖 fpdf2，缺失则静默跳过）。

    v3.45.4（EXPORT-LINK）：改用 clone_from 克隆整册（**含书签 /Outlines**）后
    逐页 merge_page 盖印。旧实现用全新 PdfWriter 逐页 add_page，整书 PDF 的
    层级书签会被全部丢掉（合并时写入 3 条、盖印后 outline 变空）——
    这正是「转化为 PDF 缺少书签」的根因。

    返回是否成功盖印。"""
    try:
        from fpdf import FPDF
        from pypdf import PdfReader, PdfWriter
        import io
    except ImportError:
        _log.warning("缺少 fpdf2，跳过页码盖印")
        return False
    try:
        reader = PdfReader(pdf_path)
        n = len(reader.pages)
        if n == 0:
            return False
        mb = reader.pages[0].mediabox
        w = float(mb.width)
        h = float(mb.height)
        pdf = FPDF(unit="pt", format=(w, h))
        pdf.set_auto_page_break(False)
        for i in range(start, start + n):
            pdf.add_page()
            pdf.set_font("Helvetica", size=size)
            pdf.set_text_color(*color)
            pdf.set_xy(0, h - 20)
            pdf.cell(w, 0, str(i), align="C")
        buf = io.BytesIO()
        pdf.output(buf)
        overlay = PdfReader(io.BytesIO(buf.getvalue()))
        writer = PdfWriter(clone_from=pdf_path)
        for i, page in enumerate(writer.pages):
            page.merge_page(overlay.pages[i])
        with open(pdf_path, "wb") as f:
            writer.write(f)
        return True
    except Exception as e:
        _log.warning("页码盖印失败（已忽略）: %s", e)
        return False


def _uri_to_local_path(uri):
    """file:// URI → 本地路径（非 file scheme / 解析失败返回 None）。"""
    try:
        from urllib.parse import urlparse, unquote
        u = urlparse(str(uri))
        if u.scheme.lower() != "file":
            return None
        p = unquote(u.path or "")
        if re.match(r"^/[A-Za-z]:", p):     # Windows：/C:/x → C:/x
            p = p[1:]
        return p
    except Exception:
        return None


def _page_lookup(path_to_page, local):
    """按本地路径查章节页号；.md/.markdown/.mdown/.html/.htm 扩展名互换后再查一次。

    v3.45.6（PDF-LINK-3）：导出 HTML 的 href 既可能是 .md（getExportHTML 的
    relFromDocDir）也可能是 .html（某些改写路径），Chromium 注解 URI 原样
    带出；旧实现只按原样 key 查，.html 形态命中不了 .md 的 page_map →
    落入相对 .html 死链分支（书籍 PDF 并不产出 .html 文件 → 点了没反应）。
    """
    if not path_to_page or not local:
        return None
    key = os.path.normcase(os.path.normpath(local)).lower()
    pg = path_to_page.get(key)
    if pg is not None:
        return pg
    root, ext = os.path.splitext(local)
    el = ext.lower()
    if el in (".md", ".markdown", ".mdown", ".html", ".htm"):
        for alt_ext in (".md", ".html", ".markdown", ".mdown", ".htm"):
            if alt_ext == el:
                continue
            alt = os.path.normcase(os.path.normpath(root + alt_ext)).lower()
            pg = path_to_page.get(alt)
            if pg is not None:
                return pg
    return None


def linkify_pdf_links(pdf_path, path_to_page=None, out_dir=None, base_dir=None):
    """把 PDF 里的本地文件链接注解改写为可点击目标。

    v3.45.3（EXPORT-LINK）：命中本书章节（path_to_page）→ PDF 内部页跳转。
    v3.45.6（PDF-LINK-3）：GoTo 目标由 add_named_destination 的「目标字典」
    间接引用改为**显式目标数组** `[page_ref /FitH 页高]`（ISO 32000 只认
    名字/字符串/数组，字典形态严格阅读器不跳转）；页查找兼容 .md↔.html
    扩展名互换（见 _page_lookup）。
    v3.45.5（PDF-LINK-2）：未命中章节时（单文档导出 / 章外链接），参考 HTML
    导出口径把本地 .md 目标改写为「相对 .html 路径」—— 多数 PDF 阅读器禁用
    file:// 协议，绝对路径链接点击无响应（「转化为 PDF 后链接打不开」的
    主因）；相对 .html 与 HTML 导出物配套（导出 HTML 时同步把 .md 链接
    改写为 .html，见 _rewrite_export_link）。

    path_to_page: {abs_path（normcase+lower 归一）: 0 基页号}，可为 None/{}。
    out_dir: 相对路径基准（PDF 所在目录）；缺省且无 path_to_page 时直接返回。
    base_dir: 解析「无协议相对 URI」注解时的基准（默认取 out_dir）。
    返回改写条数。"""
    path_to_page = path_to_page or {}
    if not path_to_page and not out_dir:
        return 0
    if base_dir is None:
        base_dir = out_dir
    try:
        from pypdf import PdfWriter
        from pypdf.generic import (NameObject, DictionaryObject,
                                   TextStringObject, ArrayObject, FloatObject)
    except ImportError:
        return 0
    n = 0
    try:
        writer = PdfWriter(clone_from=pdf_path)
        total = len(writer.pages)
        for page in writer.pages:
            annots = page.get("/Annots")
            if not annots:
                continue
            for ref in annots:
                try:
                    annot = ref.get_object()
                    act = annot.get("/A")
                    if not act:
                        continue
                    uri = act.get("/URI")
                    if not uri:
                        continue
                    uri = str(uri)
                    base_uri, sep, frag = uri.partition("#")   # 去 fragment 再匹配
                    local = _uri_to_local_path(base_uri)
                    if not local and base_dir and "://" not in base_uri \
                            and not base_uri.startswith("/"):
                        # 无协议相对 URI：按 base_dir 解析后再统一处理
                        local = os.path.normpath(
                            os.path.join(base_dir, base_uri.replace("/", os.sep)))
                    if not local:
                        continue
                    pg = _page_lookup(path_to_page, local)
                    if pg is not None and 0 <= pg < total:
                        # v3.45.6（PDF-LINK-3）：GoTo 目标改「显式目标数组」
                        # [page_ref /FitH 页高] —— 旧实现用
                        # add_named_destination 返回的「目标字典」间接引用塞进
                        # /A /D，而 ISO 32000 规定动作的 /D 只接受 名字/字符串/
                        # 数组；字典形态仅部分阅读器容忍（严格解析器含部分
                        # pdf.js 版本直接不跳转 —— 「整书 PDF 链接点了没反应」）。
                        # 显式数组是所有阅读器都支持的基础形态，且 /FitH 置
                        # 目标页页顶（页高取自该页 mediabox，退化时 826）。
                        try:
                            top = float(writer.pages[pg].mediabox.height)
                        except Exception:
                            top = 826.0
                        # 注意：pypdf 的 NumberObject 继承 int，传 float 会被
                        # 截断（840.96→840）；页高必须用 FloatObject。
                        annot[NameObject("/A")] = DictionaryObject({
                            NameObject("/S"): NameObject("/GoTo"),
                            NameObject("/D"): ArrayObject([
                                writer.pages[pg].indirect_reference,
                                NameObject("/FitH"), FloatObject(top)]),
                        })
                        n += 1
                        continue
                    if not out_dir:
                        continue
                    # v3.45.5（PDF-LINK-2）：本地 .md → 相对 .html（保留 #fragment）
                    root, ext = os.path.splitext(local)
                    if ext.lower() in (".md", ".markdown", ".mdown"):
                        local = root + ".html"
                    try:
                        rel = os.path.relpath(local, out_dir)
                    except ValueError:
                        continue   # 跨盘符等：保持原样
                    from urllib.parse import quote
                    new_uri = quote(rel.replace(os.sep, "/"), safe="/") \
                        + (sep + frag if sep else "")
                    annot[NameObject("/A")] = DictionaryObject({
                        NameObject("/S"): NameObject("/URI"),
                        NameObject("/URI"): TextStringObject(new_uri),
                    })
                    n += 1
                except Exception as e:   # noqa: BLE001  单条失败不影响其余
                    _log.debug("链接注解改写跳过：%s", e)
                    continue
        if n:
            with open(pdf_path, "wb") as f:
                writer.write(f)
        return n
    except Exception as e:
        _log.warning("PDF 内链改写失败（已忽略）: %s", e)
        return 0


# ---------------------------------------------------------------- pandoc 通道

def find_pandoc(configured=""):
    """返回可用的 pandoc 可执行文件路径（未找到返回空串）。"""
    if configured:
        p = configured.strip().strip('"')
        if os.path.isfile(p):
            return p
        w = shutil.which(p)
        if w:
            return w
        return ""
    return shutil.which("pandoc") or ""


def pandoc_convert(pandoc, md, out_path, to, base_dir=""):
    """调用外部 pandoc 转换（docx / epub）。返回 (是否成功, 说明)。"""
    import tempfile
    tmp = None
    try:
        fd, tmp = tempfile.mkstemp(suffix=".md", text=True)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(md or "")
        cmd = [pandoc, tmp, "-o", out_path, "--from", "gfm+tex_math_dollars",
               "--resource-path", base_dir or ".", "--standalone"]
        if to == "epub":
            cmd += ["--toc", "--toc-depth=3"]
        else:
            cmd += ["--toc", "--toc-depth=3"]
        r = subprocess.run(cmd, capture_output=True, timeout=180)
        if r.returncode != 0:
            return False, (r.stderr or b"").decode("utf-8", "replace")[-300:]
        return True, f"已用 pandoc 导出 · {out_path}"
    except Exception as e:
        return False, f"pandoc 调用失败：{e}"
    finally:
        if tmp and os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
