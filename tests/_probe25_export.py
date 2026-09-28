# -*- coding: utf-8 -*-
"""探针 25：导出后端（docx / EPUB / PDF 书签）独立自测，不启动 GUI。"""
import json
import os
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(os.path.dirname(HERE), "src")
sys.path.insert(0, SRC)

import exporters  # noqa: E402

MD = """# 第一章 概述

这是**粗体**、*斜体*、`代码` 与 [链接](https://example.com) 的段落。

## 1.1 小节

- 无序项 A
- 无序项 B

1. 有序项 1
2. 有序项 2

> 这是一段引用。

| 列A | 列B |
| --- | --- |
| a1 | b1 |
| a2 | b2 |

```python
print("hello")
```

```mermaid
flowchart TD
  A --> B
```

```math
E = mc^2
```

![示意图](pic.png)

---

# 第二章 结论

结束语。
"""

results = []


def check(name, ok, extra=""):
    results.append(("PASS" if ok else "FAIL", name, str(extra)[:160]))


tmp = tempfile.mkdtemp(prefix="mdexport_")
# 造一张真图片（1x1 PNG）
import base64
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
with open(os.path.join(tmp, "pic.png"), "wb") as f:
    f.write(PNG)

# ---------- 1) 块解析 ----------
blocks = exporters.parse_blocks(MD)
kinds = [b["kind"] for b in blocks]
check("parse_blocks 覆盖 heading/para/list/quote/table/fence/hr",
      kinds.count("heading") == 3 and "table" in kinds and "quote" in kinds
      and kinds.count("fence") == 3 and "hr" in kinds
      and kinds.count("list") == 2, kinds)
fences = [b["lang"] for b in blocks if b["kind"] == "fence"]
check("围栏语言识别 python/mermaid/math", fences == ["python", "mermaid", "math"], fences)

# ---------- 2) 行内解析 ----------
inl = exporters.parse_inline("a **b** *c* `d` ~~e~~ [f](g) ![h](i.png)")
props = {t.get("text"): t for t in inl if t.get("text")}
check("行内：粗体/斜体/代码/删除线/链接",
      props.get("b", {}).get("bold") and props.get("c", {}).get("italic")
      and props.get("d", {}).get("code") and props.get("e", {}).get("strike")
      and props.get("f", {}).get("link") == "g", inl)

# ---------- 3) docx ----------
docx_path = os.path.join(tmp, "out.docx")
ok, msg = exporters.write_docx(MD, docx_path, tmp, "测试文档")
check("write_docx 成功", ok and os.path.getsize(docx_path) > 5000, msg)
if ok:
    import docx
    d = docx.Document(docx_path)
    n_par = len(d.paragraphs)
    texts = "\n".join(p.text for p in d.paragraphs)
    check("docx 含标题/表格/代码", n_par > 15 and len(d.tables) == 1, f"par={n_par} tbl={len(d.tables)}")
    check("docx 正文文字完整", "第一章 概述" in texts and "结束语" in texts and 'print("hello")' in texts,
          texts[:80])
    check("docx 图表/公式降级为标注", "Mermaid" in texts and "LaTeX" in texts)
    check("docx 图片已嵌入",
          any(r.__class__.__name__ == "InlineShape" or True for r in d.inline_shapes)
          and len(d.inline_shapes) >= 1, f"shapes={len(d.inline_shapes)}")

# ---------- 4) EPUB ----------
epub_path = os.path.join(tmp, "out.epub")
ok2, msg2 = exporters.write_epub(MD, epub_path, tmp, "测试文档")
check("write_epub 成功", ok2 and os.path.getsize(epub_path) > 1000, msg2)
if ok2:
    z = zipfile.ZipFile(epub_path)
    names = z.namelist()
    check("epub 结构完整（mimetype/container/opf/nav/css）",
          "mimetype" in names and "META-INF/container.xml" in names
          and "OEBPS/content.opf" in names and "OEBPS/nav.xhtml" in names
          and "OEBPS/style.css" in names, names[:8])
    zi = z.getinfo("mimetype")
    check("epub mimetype 未压缩且内容正确",
          zi.compress_type == zipfile.ZIP_STORED
          and z.read("mimetype").decode() == "application/epub+zip")
    check("epub 按一级标题分 2 章",
          "OEBPS/chapter1.xhtml" in names and "OEBPS/chapter2.xhtml" in names, names)
    ch1 = z.read("OEBPS/chapter1.xhtml").decode("utf-8")
    check("epub 章节含表格/图片/代码",
          "<table>" in ch1 and "<img" in ch1 and "<pre>" in ch1 and "images/pic.png" in ch1,
          ch1[:120])
    check("epub 图片已打包", "OEBPS/images/pic.png" in names, [n for n in names if "images" in n])
    opf = z.read("OEBPS/content.opf").decode("utf-8")
    check("epub opf 含标题与 manifest 图片", "测试文档" in opf and "images/pic.png" in opf)

# ---------- 5) 无图片/空文档鲁棒性 ----------
ok3, msg3 = exporters.write_docx("", os.path.join(tmp, "empty.docx"), tmp)
check("空文档导出 docx 不报错", ok3, msg3)
ok4, msg4 = exporters.write_epub("# 只有标题\n", os.path.join(tmp, "one.epub"), tmp)
check("单章 epub 正常", ok4, msg4)

# ---------- 6) PDF 书签 ----------
try:
    from pypdf import PdfWriter
    pdf_path = os.path.join(tmp, "bk.pdf")
    w = PdfWriter()
    for _ in range(3):
        w.add_blank_page(width=595, height=842)
    with open(pdf_path, "wb") as f:
        w.write(f)
    n, msg5 = exporters.add_pdf_bookmarks(
        pdf_path, [{"level": 1, "text": "第一章", "index": 0}])
    check("pypdf 可用（无匹配文本时安全降级）", n == 0 and "未" in msg5, msg5)
except ImportError:
    check("pypdf 可用", False, "未安装 pypdf")

# ---------- 7) pandoc 探测 ----------
p = exporters.find_pandoc("")
check("find_pandoc 在未安装时返回空串（不抛异常）", p == "", repr(p))
check("find_pandoc 接受无效路径", exporters.find_pandoc("D:/nope/pandoc.exe") == "")

print(json.dumps({"results": results}, ensure_ascii=False))
with open(os.path.join(HERE, "_probe25_export_out.json"), "w", encoding="utf-8") as f:
    json.dump(results, f, ensure_ascii=False, indent=1)
print("PROBE25 DONE")
