# -*- coding: utf-8 -*-
"""v3.30.19 批次 C：exporters.py（MATH-6 识别 / MMD-2 PNG 嵌入）+ mainwindow.py（导出链路 / 图表主题菜单）"""
import io, sys

fails = []
applied = []

def patch(path, pairs):
    src = io.open(path, encoding="utf-8").read()
    for tag, old, new in pairs:
        n = src.count(old)
        if n != 1:
            fails.append("%s @ %s: expect 1 hit, got %d" % (tag, path, n))
            continue
        src = src.replace(old, new)
        applied.append(tag)
    io.open(path, "w", encoding="utf-8", newline="").write(src)

# ================= exporters.py =================
EXP = r"E:\WorkBuddy\md-editor\src\exporters.py"
patch(EXP, [
    # ---- MATH-6：$$ 块级公式识别（按公式源码块降级导出，与 ```math 同路径）----
    ("exp-math6", """        if _HR.match(line):
            blocks.append({"kind": "hr"})
            i += 1
            continue
""",
"""        if _HR.match(line):
            blocks.append({"kind": "hr"})
            i += 1
            continue

        # v3.30.19（MATH-6）：$$…$$ 块级公式 → 按公式源码块降级导出（渲染由
        # 编辑器 / pandoc 通道负责；内置导出与 ```math 同路径）
        m6 = re.match(r"^\\s{0,3}\\$\\$(.*)$", line)
        if m6:
            first6 = m6.group(1)
            i += 1
            if first6.strip().endswith("$$"):
                blocks.append({"kind": "fence", "lang": "math",
                               "text": first6.strip()[:-2].strip()})
                continue
            buf6 = []
            if first6.strip():
                buf6.append(first6)
            while i < n and "$$" not in lines[i]:
                buf6.append(lines[i])
                i += 1
            if i < n:
                buf6.append(re.sub(r"\\$\\$.*$", "", lines[i]))
                i += 1
            blocks.append({"kind": "fence", "lang": "math",
                           "text": "\\n".join(buf6).strip()})
            continue
"""),
    # ---- MMD-2：write_docx 增加 mmd_png 参数 + 图表 PNG 嵌入 ----
    ("exp-docx-import", """    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError:""",
"""    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:"""),
    ("exp-docx-sig", """def write_docx(md, out_path, base_dir="", title=None):
    \"\"\"Markdown → .docx。返回 (是否成功, 提示信息)。\"\"\"""",
"""def write_docx(md, out_path, base_dir="", title=None, mmd_png=None):
    \"\"\"Markdown → .docx。返回 (是否成功, 提示信息)。

    v3.30.19（MMD-2）：mmd_png 为 {图表源码: PNG 路径} 映射（由 GUI 侧把已渲染
    SVG 位图化后传入）；命中时图表直接嵌 PNG，未命中回落「源码 + 标注」降级。
    \"\"\""""),
    ("exp-docx-mmd", """        elif k == "fence":
            lang, code = blk["lang"], blk["text"]
            if lang == "mermaid":
                tip = doc.add_paragraph()""",
"""        elif k == "fence":
            lang, code = blk["lang"], blk["text"]
            if lang == "mermaid" and mmd_png and mmd_png.get(code) \\
                    and os.path.isfile(mmd_png[code]):
                # v3.30.19（MMD-2）：已位图化的图表直接嵌 PNG（2x 渲染，白底）
                try:
                    p = doc.add_paragraph()
                    p.alignment = 1  # 居中
                    p.add_run().add_picture(mmd_png[code], width=Inches(5.6))
                    continue
                except Exception as e2:
                    _log.debug("docx 图表 PNG 嵌入失败，回落源码降级: %s", e2)
            if lang == "mermaid":
                tip = doc.add_paragraph()"""),
    # ---- MMD-2：EPUB 同链路 ----
    ("exp-epub-sig", """def _blocks_to_xhtml(blocks, base_dir, images):
    \"\"\"块列表 → XHTML 片段；顺带收集需要打包进 EPUB 的图片。\"\"\"""",
"""def _blocks_to_xhtml(blocks, base_dir, images, mmd_png=None):
    \"\"\"块列表 → XHTML 片段；顺带收集需要打包进 EPUB 的图片。

    v3.30.19（MMD-2）：mmd_png 命中的图表直接以 PNG 内嵌（随 EPUB 打包）。
    \"\"\""""),
    ("exp-epub-mmd", """        elif k == "fence":
            lang, code = blk["lang"], blk["text"] or ""
            if lang in ("mermaid", "math"):
                label = "图表（Mermaid 源码）" if lang == "mermaid" else "公式（LaTeX 源码）"
                out.append(f'<p class="tip">［{label}］</p>')
            out.append(f"<pre><code>{_xesc(code)}</code></pre>")""",
"""        elif k == "fence":
            lang, code = blk["lang"], blk["text"] or ""
            if lang == "mermaid" and mmd_png and mmd_png.get(code) \\
                    and os.path.isfile(mmd_png[code]):
                # v3.30.19（MMD-2）：图表 PNG 随 EPUB 打包内嵌
                img_name = f"mmd-{len(images) + 1}.png"
                images[img_name] = mmd_png[code]
                out.append(f'<p><img src="images/{img_name}" alt="图表"/></p>')
                continue
            if lang in ("mermaid", "math"):
                label = "图表（Mermaid 源码）" if lang == "mermaid" else "公式（LaTeX 源码）"
                out.append(f'<p class="tip">［{label}］</p>')
            out.append(f"<pre><code>{_xesc(code)}</code></pre>")"""),
    ("exp-epub-sig2", """def write_epub(md, out_path, base_dir="", title=None, author="mdEditor"):
    \"\"\"Markdown → EPUB 3。返回 (是否成功, 提示信息)。\"\"\"""",
"""def write_epub(md, out_path, base_dir="", title=None, author="mdEditor", mmd_png=None):
    \"\"\"Markdown → EPUB 3。返回 (是否成功, 提示信息)。（MMD-2：mmd_png 见 write_docx）\"\"\""""),
    ("exp-epub-call", """        body = _blocks_to_xhtml(ch, base_dir, images)""",
"""        body = _blocks_to_xhtml(ch, base_dir, images, mmd_png=mmd_png)"""),
])

# ================= mainwindow.py =================
MW = r"E:\WorkBuddy\md-editor\src\mainwindow.py"
patch(MW, [
    # ---- MMD-7：图表主题随视图加载同步 ----
    ("mw-mmd7-load", """        self._js_view(view, f"window.setCodePrefs({{lnum: {str(self._code_lnum).lower()}, "
                            f"style: {js(self._code_style)}}})")""",
"""        self._js_view(view, f"window.setCodePrefs({{lnum: {str(self._code_lnum).lower()}, "
                            f"style: {js(self._code_style)}}})")
        # v3.30.19（MMD-7）：全局图表主题随视图加载同步（空 = 跟随界面主题）
        self._js_view(view, f"window.setMmdTheme({js(self._settings.value('mmdTheme', '') or '')})")"""),
    # ---- MMD-7：视图菜单「图表主题」子菜单 ----
    ("mw-mmd7-menu", """        m_view.addSeparator()
        self._reg_act("视图", "全屏", "F11", self.toggle_fullscreen, menu=m_view)""",
"""        m_view.addSeparator()
        # v3.30.19（MMD-7）：全局图表主题（mermaid 主题；「跟随界面」= 空 = 自动）
        m_mmd = m_view.addMenu("图表主题")
        self._mmd_group = QActionGroup(self)
        self._mmd_group.setExclusive(True)
        for label, val in (("跟随界面主题", ""), ("默认（浅色）", "default"),
                           ("深色", "dark"), ("森林", "forest"), ("中性", "neutral")):
            a = QAction(label, self)
            a.setCheckable(True)
            a.setChecked((self._settings.value("mmdTheme", "") or "") == val)
            a.triggered.connect(lambda _c=False, v=val: self._set_mmd_theme(v))
            self._mmd_group.addAction(a)
            m_mmd.addAction(a)
        m_view.addSeparator()
        self._reg_act("视图", "全屏", "F11", self.toggle_fullscreen, menu=m_view)"""),
    # ---- MMD-7：主题切换落盘 + 下发 ----
    ("mw-mmd7-set", """    def _reg_act(self, group, name, seq, slot, menu=None, checkable=False):""",
"""    def _set_mmd_theme(self, val):
        \"\"\"v3.30.19（MMD-7）：全局图表主题切换（持久化 + 下发全部已加载视图）。\"\"\"
        self._settings.setValue("mmdTheme", val)
        self._js_all(f"window.setMmdTheme({js(val)})")
        self.show_status("图表主题已更新" + ("" if val else "（跟随界面主题）"))

    def _js_all(self, code):
        \"\"\"对全部已加载的编辑器视图执行 JS。\"\"\"
        for role, view in (("edit", self._view), ("src", getattr(self, "_view_src", None)),
                           ("prev", getattr(self, "_view_prev", None))):
            st = self._vstate.get(role)
            if view is not None and st and st.get("loaded"):
                self._js_view(view, code)

    def _reg_act(self, group, name, seq, slot, menu=None, checkable=False):"""),
    # ---- MMD-2：导出链路（图表位图化后交内置导出器） ----
    ("mw-mmd2-export", """            path, _ = QFileDialog.getSaveFileName(
                self, f"导出 {label}",
                self._dlg_start(kind,
                    os.path.splitext(self.file_path or "未命名")[0] + ext), filt)
            if not path:
                return
            self._remember_dir(kind, path)
            base_dir = self._doc_dir()
            title = os.path.splitext(os.path.basename(self.file_path or "未命名"))[0]
            if pandoc:
                ok, msg = exporters.pandoc_convert(pandoc, md, path, kind, base_dir)
            elif kind == "docx":
                ok, msg = exporters.write_docx(md, path, base_dir, title)
            else:
                ok, msg = exporters.write_epub(md, path, base_dir, title)
            self.show_status(msg)
            if not ok:
                QMessageBox.warning(self, f"导出 {label} 失败", msg)

        self._js_active("window.getMarkdown()", cb)""",
"""            path, _ = QFileDialog.getSaveFileName(
                self, f"导出 {label}",
                self._dlg_start(kind,
                    os.path.splitext(self.file_path or "未命名")[0] + ext), filt)
            if not path:
                return
            self._remember_dir(kind, path)
            base_dir = self._doc_dir()
            title = os.path.splitext(os.path.basename(self.file_path or "未命名"))[0]
            if pandoc:
                ok, msg = exporters.pandoc_convert(pandoc, md, path, kind, base_dir)
                self.show_status(msg)
                if not ok:
                    QMessageBox.warning(self, f"导出 {label} 失败", msg)
                return
            # v3.30.19（MMD-2）：内置导出前先把已渲染图表位图化（SVG→PNG）
            self._export_with_mmd(kind, label, md, path, base_dir, title)

        self._js_active("window.getMarkdown()", cb)

    def _export_with_mmd(self, kind, label, md, path, base_dir, title):
        \"\"\"v3.30.19（MMD-2）：让编辑器把已渲染 SVG 转 PNG，收齐后执行内置导出。\"\"\"
        def start(total):
            if not total:
                self._export_now(kind, label, md, path, base_dir, title, {})
                return
            self._mmd_export = {"total": int(total), "t0": time.time()}
            self._mmd_poll(kind, label, md, path, base_dir, title)
        self._js_active("window.__mmdExportStart ? window.__mmdExportStart() : 0", start)

    def _mmd_poll(self, kind, label, md, path, base_dir, title):
        def got(state):
            st = state or {}
            done = int(st.get("done", 0) or 0)
            total = int(st.get("total", 0) or 0)
            if done >= total or time.time() - self._mmd_export["t0"] > 12:
                mmd_png = {}
                for it in st.get("items", []):
                    src = str(it.get("src") or "")
                    data = str(it.get("data") or "")
                    if not src or not data or src in mmd_png:
                        continue
                    try:
                        fd, tmp = tempfile.mkstemp(suffix=".png")
                        with os.fdopen(fd, "wb") as f:
                            f.write(base64.b64decode(data.split(",", 1)[-1]))
                        mmd_png[src] = tmp
                    except Exception as e3:
                        _log.debug("图表 PNG 落盘失败: %s", e3)
                self._export_now(kind, label, md, path, base_dir, title, mmd_png)
                return
            QTimer.singleShot(300, lambda: self._mmd_poll(kind, label, md, path, base_dir, title))
        self._js_active("window.__mmdExportState || null", got)

    def _export_now(self, kind, label, md, path, base_dir, title, mmd_png):
        if kind == "docx":
            ok, msg = exporters.write_docx(md, path, base_dir, title, mmd_png=mmd_png)
        else:
            ok, msg = exporters.write_epub(md, path, base_dir, title, mmd_png=mmd_png)
        self.show_status(msg)
        if not ok:
            QMessageBox.warning(self, f"导出 {label} 失败", msg)"""),
    # ---- MMD-2：tempfile import ----
    ("mw-import", """import subprocess
import time""",
"""import subprocess
import tempfile
import time"""),
])

if fails:
    print("FAILED:", len(fails))
    for f in fails: print("  -", f)
    sys.exit(1)
print("ALL OK,", len(applied), "edits")
